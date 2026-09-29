#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsxtomysql 的 Excel 读写助手。

这个文件的全部内容会被逐字内嵌进 xlsxtomysql-rs 的 Rust 源码常量 PY_HELPER，
由宿主以 `python3 -c <本源码> <mode> [参数...]` 的方式调用 —— 运行时不落地成文件。

协议
----
stdout：制表符分隔的文本行。以 # 开头的是控制行，其余为数据行。
    #KIND<TAB>格式名                    xlsx / Excel 2003 (.xls)
    #SHEET<TAB>表名<TAB>总行<TAB>总列
    #SHEETS<TAB>表1<TAB>表2...
    #MERGES<TAB>r1,c1,r2,c2|r1,c1,r2,c2|...
    #FORMULA<TAB>0|1
    #HEADER<TAB>单元格...
    #ROW<TAB>行号<TAB>单元格...
    #END<TAB>读取行数<TAB>最后读到的行号<TAB>已输出行数
    #OK<TAB>绝对路径
    #FATAL<TAB>退出码<TAB>消息        退出码与宿主进程的退出码约定一致
每个单元格编码为 "<kind>:<已转义的值>"，kind 为单字符：
    e 空 / i 整数 / d 小数 / f 科学计数法 / b 布尔 / s 文本
    D 日期 / T 时间 / M 日期时间 / x Excel 错误值
日期时间的值用规范文本：#D:2024-03-05 / T:09:30:00 / M:2024-03-05 09:30:00

stderr：#P<TAB>当前<TAB>总数  是进度，由宿主渲染成进度条；其余行原样转发给用户。
"""

import datetime as dt
import math
import os
import re
import sys
import zipfile

# --------------------------------------------------------------------------
# 协议输出
# --------------------------------------------------------------------------
TAB = "\t"


def w(line):
    sys.stdout.write(line)
    sys.stdout.write("\n")


def progress(n, total):
    sys.stderr.write("#P\t%d\t%d\n" % (n, total))
    sys.stderr.flush()


def note(msg):
    sys.stderr.write(str(msg) + "\n")
    sys.stderr.flush()


_ESC = {"\\": "\\\\", "\t": "\\t", "\n": "\\n", "\r": "\\r"}
_UNESC = {"t": "\t", "n": "\n", "r": "\r", "\\": "\\"}


def esc(s):
    out = []
    for ch in s:
        out.append(_ESC.get(ch, ch))
    return "".join(out)


def unesc(s):
    if "\\" not in s:
        return s
    out = []
    i = 0
    n = len(s)
    while i < n:
        c = s[i]
        if c == "\\" and i + 1 < n:
            out.append(_UNESC.get(s[i + 1], s[i + 1]))
            i += 2
        else:
            out.append(c)
            i += 1
    return "".join(out)


# --------------------------------------------------------------------------
# 值分类（与 Python 版 xlsxtomysql.py 保持同一套判定规则）
# --------------------------------------------------------------------------
K_EMPTY, K_INT, K_DEC, K_FLOAT = "e", "i", "d", "f"
K_BOOL, K_TEXT = "b", "s"
K_DATE, K_TIME, K_DATETIME = "D", "T", "M"
K_ERROR = "x"
K_OOR = "o"          # 日期/时间超出 MySQL 可表示范围，按文本处理

RE_INT = re.compile(r"^[+-]?\d+$")
RE_DEC = re.compile(r"^[+-]?(?:\d+\.\d*|\.\d+)$")
RE_SCI = re.compile(r"^[+-]?\d+(?:\.\d+)?[eE][+-]?\d+$")
RE_LEADING_ZERO = re.compile(r"^[+-]?0\d")
RE_CONTROL = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f]")

# 与 xlsxtomysql.py 保持完全一致：先试更具体的（带毫秒/时区/秒），再试宽松的
DATE_FORMATS_DATETIME = (
    "%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%S.%f",
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S.%f",
    "%Y-%m-%d %H:%M:%S%z", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S",
    "%Y.%m.%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M",
    "%Y年%m月%d日 %H:%M:%S", "%Y年%m月%d日 %H:%M",
    "%Y-%m-%d %I:%M:%S %p", "%Y-%m-%d %I:%M %p",
    "%Y/%m/%d %I:%M:%S %p", "%Y/%m/%d %I:%M %p",
)
DATE_FORMATS_DATE = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y年%m月%d")
DATE_FORMATS_TIME = ("%H:%M:%S.%f", "%H:%M:%S", "%H:%M",
                     "%I:%M:%S %p", "%I:%M %p", "%I:%M:%S%p", "%I:%M%p")
RE_DURATION = re.compile(r"^([+-]?)(\d{1,3}):([0-5]\d):([0-5]\d)$")
RE_COMPACT_DATE = re.compile(r"^(19|20)\d{6}$")
# 带时区偏移的日期时间：MySQL 的 datetime 不存时区，偏移会被丢弃
RE_TZ_OFFSET = re.compile(r"(?:Z|[+-]\d{2}:?\d{2})$")

MYSQL_DATE_MIN = dt.date(1000, 1, 1)
MYSQL_DATE_MAX = dt.date(9999, 12, 31)
MYSQL_TIME_MAX_SEC = 838 * 3600 + 59 * 60 + 59

# 字段名语义提示：命中则对单元格额外尝试紧凑日期（20240305）
DATE_NAME_HINTS = ("日期", "出生", "生日", "时间", "date", "birth", "day")


def clean(s):
    return RE_CONTROL.sub("", str(s)).strip()


def looks_int(s):
    return RE_INT.match(s) is not None


def looks_dec(s):
    return RE_DEC.match(s) is not None or looks_int(s)


def temporal_in_range(kind, value):
    """该日期/时间能否被 MySQL 的 date/datetime/time 表示。"""
    if kind in (K_DATE, K_DATETIME):
        d = value.date() if isinstance(value, dt.datetime) else value
        if not isinstance(d, dt.date):
            return True
        return MYSQL_DATE_MIN <= d <= MYSQL_DATE_MAX
    if kind == K_TIME:
        if isinstance(value, dt.timedelta):
            secs = value.total_seconds()
        elif isinstance(value, dt.time):
            secs = value.hour * 3600 + value.minute * 60 + value.second
        else:
            return True
        return -MYSQL_TIME_MAX_SEC <= secs <= MYSQL_TIME_MAX_SEC
    return True


def parse_date_like(s, allow_compact=False):
    if len(s) < 4 or len(s) > 32:
        return None
    # 含中文时必须是"年月日"这类日期写法，避免把普通中文当日期
    if re.search(r"[\u4e00-\u9fff]", s) and not re.search(r"[年月日]", s):
        return None
    for f in DATE_FORMATS_DATETIME:
        try:
            d = dt.datetime.strptime(s, f)
            if d.tzinfo is not None:
                d = d.replace(tzinfo=None)     # 时区信息 MySQL 存不下，按字面时间保存
            return K_DATETIME, d
        except ValueError:
            pass
    for f in DATE_FORMATS_DATE:
        try:
            return K_DATE, dt.datetime.strptime(s.strip(), f).date()
        except ValueError:
            pass
    for f in DATE_FORMATS_TIME:
        try:
            return K_TIME, dt.datetime.strptime(s, f).time()
        except ValueError:
            pass
    # 超过 24 小时只能写成"时长"（25:30:00 / 100:00:00），MySQL time 支持
    m = RE_DURATION.match(s)
    if m and int(m.group(2)) >= 24:
        sign = -1 if m.group(1) == "-" else 1
        td = dt.timedelta(hours=int(m.group(2)), minutes=int(m.group(3)),
                          seconds=int(m.group(4)))
        return K_TIME, td * sign
    if allow_compact and RE_COMPACT_DATE.match(s):
        try:
            return K_DATE, dt.datetime.strptime(s, "%Y%m%d").date()
        except ValueError:
            pass
    return None


def cell(kind, value=""):
    return kind, value


CELL_EMPTY = (K_EMPTY, "")


def classify_string(raw):
    s = clean(raw)
    if s == "":
        return CELL_EMPTY
    if RE_SCI.match(s):
        try:
            return (K_FLOAT, repr(float(s)))
        except ValueError:
            pass
    if looks_int(s):
        if RE_LEADING_ZERO.match(s):
            return (K_TEXT, s)                       # 前导零 → 必须当文本
        if len(s.lstrip("+-")) > 18:
            return (K_TEXT, s)                       # 超长数字 → 文本
        try:
            return (K_INT, str(int(s)))
        except ValueError:
            return (K_TEXT, s)
    if looks_dec(s):
        if RE_LEADING_ZERO.match(s.lstrip("+-")):
            return (K_TEXT, s)
        try:
            return (K_DEC, repr(float(s)))
        except ValueError:
            return (K_TEXT, s)
    d = parse_date_like(s)
    if d:
        return cell_from_date_kind(d[0], d[1])
    return (K_TEXT, s)


def cell_from_date_kind(kind, value):
    """日期值 → 行协议。超出 MySQL 范围的一律按文本（K_OOR）输出。"""
    if not temporal_in_range(kind, value):
        if kind == K_DATETIME:
            return (K_OOR, value.strftime("%Y-%m-%d %H:%M:%S"))
        if kind == K_DATE:
            return (K_OOR, value.strftime("%Y-%m-%d"))
        return (K_OOR, fmt_time(value))
    if kind == K_DATETIME:
        return (K_DATETIME, value.strftime("%Y-%m-%d %H:%M:%S"))
    if kind == K_DATE:
        return (K_DATE, value.strftime("%Y-%m-%d"))
    if kind == K_TIME:
        return (K_TIME, fmt_time(value))
    return (K_TEXT, clean(value))


def fmt_time(t):
    if isinstance(t, dt.timedelta):
        secs = int(t.total_seconds())
    elif isinstance(t, dt.time):
        secs = t.hour * 3600 + t.minute * 60 + t.second
    else:
        secs = int(t)
    # 时长要保留完整小时数（MySQL time 支持到 838:59:59），不能对 24 取模
    sign = "-" if secs < 0 else ""
    secs = abs(secs)
    h, rem = divmod(secs, 3600)
    m, s = divmod(rem, 60)
    return "%s%02d:%02d:%02d" % (sign, h, m, s)


def classify_native_date(value, number_format=None):
    fmt = (number_format or "").lower()
    has_hour = ("h" in fmt) or ("[h]" in fmt)
    has_date = bool(re.search(r"[yd]", fmt))
    if isinstance(value, dt.datetime):
        if value.year <= 1900 and ((value.month == 1 and value.day <= 1) or value.year == 1899):
            return CELL_EMPTY
        if has_hour and not has_date:
            return cell_from_date_kind(K_TIME, value.time())
        if has_date and not has_hour:
            return cell_from_date_kind(K_DATE, value.date())
        if not has_date and not has_hour:
            return cell_from_date_kind(K_DATETIME, value)
        if value.hour == value.minute == value.second == 0 and not has_hour:
            return cell_from_date_kind(K_DATE, value.date())
        return cell_from_date_kind(K_DATETIME, value)
    if isinstance(value, dt.date):
        if value.year <= 1900:
            return CELL_EMPTY
        return cell_from_date_kind(K_DATE, value)
    if isinstance(value, dt.time):
        return cell_from_date_kind(K_TIME, value)
    if isinstance(value, dt.timedelta):
        return cell_from_date_kind(K_TIME, value)
    return (K_TEXT, clean(value))


def encode(c):
    kind, val = c
    return kind + ":" + esc(val)


# --------------------------------------------------------------------------
# 格式探测
# --------------------------------------------------------------------------
def sniff_format(path):
    with open(path, "rb") as f:
        head = f.read(8)
    if head[:2] == b"PK":
        return "xlsx"
    if head[:4] == b"\xd0\xcf\x11\xe0":
        return "xls"
    if head[:5] == b"<?xml" or head[:1] == b"<":
        return "xml"
    return "unknown"


def parse_range(ref):
    m = re.fullmatch(r"([A-Za-z]{1,3})(\d+)(?::([A-Za-z]{1,3})(\d+))?", ref.strip())
    if not m:
        return None

    def colnum(s):
        n = 0
        for ch in s.upper():
            n = n * 26 + (ord(ch) - 64)
        return n

    c1 = colnum(m.group(1))
    r1 = int(m.group(2))
    c2 = colnum(m.group(3)) if m.group(3) else c1
    r2 = int(m.group(4)) if m.group(4) else r1
    return (min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))


ROW_RE = re.compile(rb"<row[^>]*\br=\"(\d+)\"")
COL_RE = re.compile(rb"<c[^>]*\br=\"([A-Z]{1,3})\d+\"")
DIM_RE = re.compile(rb"<dimension[^>]*\bref=\"([^\"]+)\"")
MERGE_BLOCK_RE = re.compile(rb"<mergeCells[^>]*>(.*?)</mergeCells>", re.S)
MERGE_ONE_RE = re.compile(rb"<mergeCell[^>]*ref=\"([^\"]+)\"")
HCOL_RE = re.compile(rb"<col\b[^>]*hidden=\"(?:1|true)\"[^>]*/?>")
HROW_RE = re.compile(rb"<row\b[^>]*hidden=\"(?:1|true)\"[^>]*>")
ATTR_NUM_RE = re.compile(rb"(\w+)=\"(\d+)\"")

NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def xlsx_scan(path, sheet, merge_limit_mb=64, merge_scan="auto"):
    """流式扫描工作表 XML：合并区域 / 公式 / 真实行列数。"""
    import xml.etree.ElementTree as ET

    merges = []
    has_formula = False
    has_embed_img = False
    hidden_cols = []
    hidden_rows = []
    max_row = 0
    max_col = 0
    dim_rows = 0
    dim_cols = 0

    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if "xl/workbook.bin" in names:
            raise HelperError("检测到 .xlsb 二进制工作簿，openpyxl 不支持，请另存为 .xlsx 后重试")
        try:
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
            rels_xml = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        except KeyError:
            raise HelperError("不是有效的 .xlsx 文件（缺少工作簿定义）")
        rels = {r.get("Id"): r.get("Target") for r in rels_xml}
        names_in_wb = []
        target = None
        for sh in wb_xml.iter(NS_MAIN + "sheet"):
            names_in_wb.append(sh.get("name"))
            if sh.get("name") == sheet:
                target = rels.get(sh.get(NS_REL + "id"))
        if target is None:
            raise HelperError(
                "工作表 %r 不存在。可用工作表: %s" % (sheet, "、".join(names_in_wb)),
                EXIT_USAGE)
        if not target.startswith("xl/"):
            target = "xl/" + target.lstrip("/")
        target = target.replace("xl/xl/", "xl/")

        size = os.path.getsize(path)
        do_merge = merge_scan != "off" and (merge_scan == "on" or size <= merge_limit_mb * 1024 * 1024)

        tail = b""
        part = 0
        with z.open(target) as fh:
            while True:
                chunk = fh.read(1 << 20)
                if not chunk:
                    break
                part += 1
                buf = tail + chunk
                if not has_formula and (b"<f " in buf or b"<f>" in buf or b"<f/" in buf):
                    has_formula = True
                if not has_embed_img and b"DISPIMG" in buf:
                    has_embed_img = True
                for m in HROW_RE.finditer(buf):
                    for k, v in ATTR_NUM_RE.findall(m.group(0)):
                        if k == b"r":
                            r = int(v)
                            if r not in hidden_rows:
                                hidden_rows.append(r)
                for m in HCOL_RE.finditer(buf):
                    attrs = dict(ATTR_NUM_RE.findall(m.group(0)))
                    lo = int(attrs.get(b"min", b"1"))
                    hi = int(attrs.get(b"max", attrs.get(b"min", b"1")))
                    for c in range(lo, min(hi, 256) + 1):
                        if c not in hidden_cols:
                            hidden_cols.append(c)
                if part == 1:
                    m = DIM_RE.search(buf)
                    if m:
                        rng = parse_range(m.group(1).decode("ascii", "replace"))
                        if rng:
                            dim_rows, dim_cols = rng[2], rng[3]
                for m in ROW_RE.finditer(buf):
                    v = int(m.group(1))
                    if v > max_row:
                        max_row = v
                if part <= 3:
                    for m in COL_RE.finditer(buf):
                        n = 0
                        for ch in m.group(1):
                            n = n * 26 + (ch - 64)
                        if n > max_col:
                            max_col = n
                if do_merge:
                    m = MERGE_BLOCK_RE.search(buf)
                    if m:
                        for ref in MERGE_ONE_RE.findall(m.group(1)):
                            rng = parse_range(ref.decode("ascii", "replace"))
                            if rng:
                                merges.append(rng)
                                max_col = max(max_col, rng[3])
                        break
                    tail = buf[-8192:]
        max_row = max(max_row, dim_rows)
        max_col = max(max_col, dim_cols)
    return (merges, has_formula, max_row, max_col, names_in_wb,
            hidden_cols, hidden_rows, has_embed_img)


# 与宿主（Rust）约定的退出码，需要随 #FATAL 一并回报
EXIT_ERROR = 1   # 运行期错误
EXIT_USAGE = 3   # 参数/引用错误（文件不存在、工作表不存在、行号越界……）


class HelperError(Exception):
    def __init__(self, msg, code=EXIT_ERROR):
        super().__init__(msg)
        self.code = code


def list_sheets(path):
    kind = sniff_format(path)
    if kind == "xls":
        import xlrd
        book = xlrd.open_workbook(path)
        return list(book.sheet_names())
    if kind == "xlsx":
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(path) as z:
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        return [sh.get("name") for sh in wb_xml.iter(NS_MAIN + "sheet")]
    raise HelperError("无法识别的文件格式")


# --------------------------------------------------------------------------
# mode = probe
# --------------------------------------------------------------------------
def _flat_pair(c, v):
    """#HFILL 的一对（列号, 值）：值已转义，不会有真 tab，宿主按两项一组解析。"""
    return "%d\t%s" % (int(c), esc(v))


def header_fill(path, sheet, header_row, merges, kind, max_col):
    """字段名行落在纵向合并区里时，字段名在合并区左上角那一行。

    两行表头的常见写法（D1:D2 合并）会让字段名行读到空值，
    这里回读合并区上方的行，把值取出来交给宿主。
    返回 [(列号(1起), 值), ...]
    """
    rows = sorted({r1 for (r1, c1, r2, c2) in merges if r1 < header_row <= r2})
    if not rows or header_row <= 1:
        return []
    lo, hi = rows[0], min(rows[-1], header_row - 1)
    fetched = {}
    try:
        if kind == "xls":
            import xlrd
            book = xlrd.open_workbook(path)
            reader = XlsReader(book, book.sheet_by_name(sheet))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            reader = XlsxReader(wb[sheet])
        for r, cells in reader.iter_rows(lo, hi, max_col):
            fetched[r] = cells
    except Exception:
        return []
    out = []
    for (r1, c1, r2, c2) in merges:
        if not (r1 < header_row <= r2):
            continue
        row = fetched.get(r1)
        if not row:
            continue
        ci = c1 - 1
        if 0 <= ci < len(row) and row[ci][0] != K_EMPTY:
            out.append((c1, row[ci][1]))
    return out


def cmd_probe(path, sheet, header_row=0):
    header_row = int(header_row)
    if not os.path.exists(path):
        raise HelperError("文件不存在: %s" % path, EXIT_USAGE)
    kind = sniff_format(path)
    if kind == "xls":
        import xlrd
        try:
            book = xlrd.open_workbook(path, formatting_info=True)
        except Exception:
            try:
                book = xlrd.open_workbook(path)
            except Exception as e:
                raise HelperError("打开 .xls 失败（可能已加密或损坏）: %s" % e)
        names = list(book.sheet_names())
        if sheet not in names:
            raise HelperError("工作表 %r 不存在。可用工作表: %s" % (sheet, "、".join(names)),
                              EXIT_USAGE)
        sh = book.sheet_by_name(sheet)
        merges = []
        for rlo, rhi, clo, chi in getattr(sh, "merged_cells", []) or []:
            merges.append((rlo + 1, clo + 1, rhi, chi))    # xlrd 为开区间 → 闭区间
        n_rows, n_cols = int(sh.nrows), int(sh.ncols)
        if merges:
            n_cols = max(n_cols, max(m[3] for m in merges))
        hidden_cols, hidden_rows = [], []
        try:
            hidden_cols = sorted(c + 1 for c, info in (getattr(sh, "colinfo_map", None) or {}).items()
                                 if getattr(info, "hidden", 0))
            hidden_rows = sorted(r + 1 for r, info in (getattr(sh, "rowinfo_map", None) or {}).items()
                                 if getattr(info, "hidden", 0))
        except Exception:
            pass
        w("#KIND\tExcel 2003 (.xls)")
        w("#SHEET\t%s\t%d\t%d" % (esc(sheet), n_rows, n_cols))
        w("#SHEETS\t%s" % TAB.join(esc(x) for x in names))
        w("#MERGES\t%s" % "|".join("%d,%d,%d,%d" % m for m in merges))
        w("#FORMULA\t0")
        w("#HIDDEN\t%s\t%s" % (",".join(str(x) for x in hidden_cols),
                               ",".join(str(x) for x in hidden_rows)))
        w("#EMBEDIMG\t0")
        hf = header_fill(path, sheet, header_row, merges, "xls", n_cols)
        w("#HFILL\t" + TAB.join(_flat_pair(c, v) for c, v in hf))
        return 0

    if kind == "xlsx":
        (merges, has_formula, max_row, max_col, names,
         hidden_cols, hidden_rows, has_img) = xlsx_scan(path, sheet)
        w("#KIND\tExcel 2007+ (.xlsx)")
        w("#SHEET\t%s\t%d\t%d" % (esc(sheet), max_row, max_col))
        w("#SHEETS\t%s" % TAB.join(esc(x) for x in names))
        w("#MERGES\t%s" % "|".join("%d,%d,%d,%d" % m for m in merges))
        w("#FORMULA\t%d" % (1 if has_formula else 0))
        w("#HIDDEN\t%s\t%s" % (",".join(str(x) for x in hidden_cols),
                               ",".join(str(x) for x in hidden_rows)))
        w("#EMBEDIMG\t%d" % (1 if has_img else 0))
        hf = header_fill(path, sheet, header_row, merges, "xlsx", max_col)
        w("#HFILL\t" + TAB.join(_flat_pair(c, v) for c, v in hf))
        return 0

    raise HelperError("无法识别的文件格式（既不是 .xlsx 也不是 .xls），请检查文件是否完整")


# --------------------------------------------------------------------------
# mode = dump
# --------------------------------------------------------------------------
def cmd_dump(path, sheet, header_row, data_row, last_row, step, offset, hints, n_cols,
             warm=0, max_read=0):
    kind = sniff_format(path)
    date_hints = None
    reader = None

    if kind == "xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise HelperError("打开 .xlsx 失败: %s" % e)
        if sheet not in wb.sheetnames:
            raise HelperError("工作表 %r 不存在。可用工作表: %s"
                              % (sheet, "、".join(wb.sheetnames)))
        ws = wb[sheet]
        reader = XlsxReader(ws)
    elif kind == "xls":
        import xlrd
        try:
            book = xlrd.open_workbook(path)
        except Exception as e:
            raise HelperError("打开 .xls 失败: %s" % e)
        if sheet not in book.sheet_names():
            raise HelperError("工作表 %r 不存在。可用工作表: %s"
                              % (sheet, "、".join(book.sheet_names())), EXIT_USAGE)
        reader = XlsReader(book, book.sheet_by_name(sheet))
    else:
        raise HelperError("无法识别的文件格式")

    # 列数由宿主（probe 阶段的真实扫描结果）给出：流式写出的 xlsx 其
    # <dimension> 常常只写 A1，不能信 ws.max_column。
    n_cols = max(int(n_cols), 1)
    last = reader.n_rows if last_row <= 0 else last_row
    if last < header_row:
        raise HelperError("第 %d 行（字段名称所在行）超出工作表范围（共 %d 行）"
                          % (header_row, reader.n_rows), EXIT_USAGE)

    total = max(last - header_row + 1, 1)
    tick = max(total // 100, 1)
    read = 0
    last_seen = header_row - 1
    emitted = 0

    for r, cells in reader.iter_rows(header_row, last, n_cols):
        if max_read and read >= max_read:
            break
        read += 1
        last_seen = r
        if read % tick == 0 or read == total:
            progress(read, total)

        if r == header_row:
            if date_hints is None:
                names = [c[1] for c in cells]
                date_hints = [bool(hints) and any(
                    h in (nm or "").lower() for h in DATE_NAME_HINTS) for nm in names]
            w("#HEADER\t" + TAB.join(encode(c) for c in cells))
            continue

        if r < data_row:
            continue
        # 开头 warm 行无条件判定，其余按步长跳跃抽样
        if step > 1 and (r - data_row) >= warm and (r - offset) % step != 0:
            continue

        if date_hints is None:
            date_hints = [False] * n_cols
        for i in range(len(cells)):
            if i >= len(date_hints):
                break
            if date_hints[i] and cells[i][0] in (K_TEXT, K_INT, K_DEC) and cells[i][1]:
                d = parse_date_like(cells[i][1], allow_compact=True)
                if d:
                    cells[i] = cell_from_date_kind(d[0], d[1])
        w("#ROW\t%d\t%s" % (r, TAB.join(encode(c) for c in cells)))
        emitted += 1

    progress(total, total)
    pct = getattr(reader, "pct_cells", 0)
    if pct:
        w("#PCT\t%d" % pct)
    tz = getattr(reader, "tz_dropped", None) or {}
    if tz:
        w("#TZ\t" + TAB.join("%d\t%d" % (c, tz[c]) for c in sorted(tz)))
    w("#END\t%d\t%d\t%d" % (read, last_seen, emitted))
    return 0


class XlsxReader:
    """只读模式单次顺序遍历：整个工作表只解析一遍 XML，绝不按行重开流。"""

    def __init__(self, ws):
        self.ws = ws
        self.n_rows = ws.max_row or 0
        self.n_cols = ws.max_column or 0
        self.pct_cells = 0       # 带百分号格式的数值单元格
        self.tz_dropped = {}     # 列号(1起) -> 带时区偏移、偏移被丢弃的单元格数

    def iter_rows(self, first, last, n_cols):
        n_cols = max(n_cols, 1)
        r = first
        for row in self.ws.iter_rows(min_row=first, max_row=last,
                                     min_col=1, max_col=n_cols):
            if not row:
                continue
            r = getattr(row[0], "row", None) or r
            while len(row) < n_cols:
                row = tuple(row) + (None,)
            cells = []
            for ci, c in enumerate(row[:n_cols]):
                v = getattr(c, "value", None) if c is not None else None
                # 百分比格式：Excel 里显示 12.5%，实际存储值是 0.125
                if isinstance(v, (int, float)) and not isinstance(v, bool) \
                        and "%" in (getattr(c, "number_format", "") or ""):
                    self.pct_cells += 1
                t = cell_from_openpyxl(c)
                # 带时区偏移的日期时间：MySQL 的 datetime 存不下时区，偏移会被丢掉
                if t[0] in (K_DATE, K_DATETIME, K_TIME) and isinstance(v, str) \
                        and RE_TZ_OFFSET.search(v.strip()):
                    self.tz_dropped[ci + 1] = self.tz_dropped.get(ci + 1, 0) + 1
                cells.append(t)
            yield r, cells
            r += 1


def cell_from_openpyxl(c):
    v = getattr(c, "value", None) if c is not None else None
    if v is None:
        return CELL_EMPTY
    if isinstance(v, bool):
        return (K_BOOL, "1" if v else "0")
    if isinstance(v, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        return classify_native_date(v, getattr(c, "number_format", None))
    if isinstance(v, (int, float)):
        if isinstance(v, float):
            if math.isnan(v) or math.isinf(v):
                return (K_ERROR, "")
            if v.is_integer() and abs(v) < 1e15:
                return (K_INT, str(int(v)))
            return (K_DEC, repr(v))
        return (K_INT, str(v))
    if isinstance(v, str):
        if getattr(c, "data_type", None) == "e":
            return (K_ERROR, clean(v))
        return classify_string(v)
    return (K_TEXT, clean(v))


class XlsReader:
    def __init__(self, book, sh):
        self.XL = __import__("xlrd")
        self.book = book
        self.sh = sh
        self.n_rows = int(sh.nrows)
        self.n_cols = int(sh.ncols)
        self.pct_cells = 0
        self.tz_dropped = {}     # 列号(1起) -> 带时区偏移、偏移被丢弃的单元格数

    def iter_rows(self, first, last, n_cols):
        for r in range(first, last + 1):
            yield r, self.row(r, max(n_cols, 1))

    def row(self, r, n_cols):
        XL = self.XL
        datemode = self.book.datemode
        sh = self.sh
        ri = r - 1
        out = []
        for ci in range(n_cols):
            if ci >= sh.ncols or ri >= sh.nrows:
                out.append(CELL_EMPTY)
                continue
            ctype = sh.cell_type(ri, ci)
            value = sh.cell_value(ri, ci)
            if ctype in (XL.XL_CELL_EMPTY, XL.XL_CELL_BLANK):
                out.append(CELL_EMPTY)
            elif ctype == XL.XL_CELL_TEXT:
                t = classify_string(str(value))
                # 带时区偏移的日期时间：MySQL 的 datetime 存不下时区，偏移会被丢掉
                if t[0] in (K_DATE, K_DATETIME, K_TIME) and RE_TZ_OFFSET.search(str(value).strip()):
                    self.tz_dropped[ci + 1] = self.tz_dropped.get(ci + 1, 0) + 1
                out.append(t)
            elif ctype == XL.XL_CELL_BOOLEAN:
                out.append((K_BOOL, "1" if value else "0"))
            elif ctype == XL.XL_CELL_ERROR:
                out.append((K_ERROR, ""))
            elif ctype == XL.XL_CELL_DATE:
                try:
                    d = XL.xldate_as_datetime(value, datemode)
                except Exception:
                    out.append((K_TEXT, clean(value)))
                    continue
                if d.year <= 1900:
                    out.append((K_TIME, fmt_time(d.time())))
                elif d.hour == d.minute == d.second == 0:
                    out.append((K_DATE, d.strftime("%Y-%m-%d")))
                else:
                    out.append((K_DATETIME, d.strftime("%Y-%m-%d %H:%M:%S")))
            else:
                if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
                    out.append((K_INT, str(int(value))))
                else:
                    out.append((K_DEC, repr(float(value))))
        return out


# --------------------------------------------------------------------------
# 极简 xlsx 写出：只用标准库 zipfile 拼 XML，不依赖 openpyxl，也不产生任何
# 临时文件。openpyxl 保存时会建中转文件再删掉，在 /tmp 很小（10MB tmpfs）或
# 进程不允许删文件的受限环境里会直接失败；errrows.xlsx 结构简单，自己写更稳。
# --------------------------------------------------------------------------
_XLSX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>"""

_XLSX_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>"""

_XLSX_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>"""

# 多工作表时按 sheet 数量动态生成（%s 处填入各 sheet 的 Override / 关系 / 声明）
_XLSX_CT_TMPL = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>%s</Types>"""

_XLSX_RELS_TMPL = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">%s</Relationships>"""

_XLSX_WB_TMPL = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>%s</sheets></workbook>"""

# XML 1.0 不允许这些控制字符，Excel 也会拒绝加载
_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")
_CELL_MAX = 32767          # Excel 单元个字符上限
_SHEET_MAX = 31            # 工作表名长度上限


def _xml_text(v):
    s = _XML_BAD.sub("", str(v))
    if len(s) > _CELL_MAX:
        s = s[:_CELL_MAX]
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _xml_attr(v):
    return _xml_text(v).replace('"', "&quot;")


def col_name(i):
    """1 -> A, 27 -> AA"""
    name = ""
    while i > 0:
        i, r = divmod(i - 1, 26)
        name = chr(65 + r) + name
    return name


def _num_text(v):
    if isinstance(v, int):
        return str(v)
    return repr(float(v))


def _cell_xml(ref, v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return '<c r="%s" t="b"><v>%d</v></c>' % (ref, 1 if v else 0)
    if isinstance(v, (int, float)):
        return '<c r="%s"><v>%s</v></c>' % (ref, _num_text(v))
    if v == "":
        return ""
    return '<c r="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (
        ref, _xml_text(v))


def write_xlsx(path, sheets):
    """把 sheets 写成一个极简 xlsx。

    sheets: [(工作表名, rows, dims)]；rows 可迭代（每项一行的 list/tuple），
    dims=(总行数, 总列数) 可为 None。逐行流式写入，内存占用与行数无关；
    不建临时文件、不删除任何文件。dims 给了就写出 <dimension>，方便其它工具
    直接取范围（只读模式下 openpyxl 拿得到 max_row/max_column，不用扫全表）。
    工作表名的合法性与去重由宿主负责，这里只截断到 31 字符。
    """
    n = len(sheets)
    overrides = "".join(
        '<Override PartName="/xl/worksheets/sheet%d.xml" ContentType='
        '"application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' % i
        for i in range(1, n + 1))
    rels = "".join(
        '<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/worksheet" Target="worksheets/sheet%d.xml"/>' % (i, i)
        for i in range(1, n + 1))
    decls = "".join(
        '<sheet name="%s" sheetId="%d" r:id="rId%d"/>'
        % (_xml_attr((nm or ("Sheet%d" % i))[:_SHEET_MAX]) or ("Sheet%d" % i), i, i)
        for i, (nm, _rows, _dims) in enumerate(sheets, 1))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _XLSX_CT_TMPL % overrides)
        z.writestr("_rels/.rels", _XLSX_ROOT_RELS)
        z.writestr("xl/workbook.xml", _XLSX_WB_TMPL % decls)
        z.writestr("xl/_rels/workbook.xml.rels", _XLSX_RELS_TMPL % rels)
        for i, (_nm, rows, dims) in enumerate(sheets, 1):
            dim = ""
            if dims and dims[0] > 0 and dims[1] > 0:
                dim = '<dimension ref="A1:%s%d"/>' % (col_name(dims[1]), dims[0])
            with z.open("xl/worksheets/sheet%d.xml" % i, "w") as f:
                f.write(b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        b'<worksheet xmlns="http://schemas.openxmlformats.org/'
                        b'spreadsheetml/2006/main">')
                if dim:
                    f.write(dim.encode("utf-8"))
                f.write(b"<sheetData>")
                for ri, row in enumerate(rows, 1):
                    f.write(('<row r="%d">' % ri).encode("utf-8"))
                    for ci, v in enumerate(row, 1):
                        xml = _cell_xml("%s%d" % (col_name(ci), ri), v)
                        if xml:
                            f.write(xml.encode("utf-8"))
                    f.write(b"</row>")
                f.write(b"</sheetData></worksheet>")


# --------------------------------------------------------------------------
# mode = errrows   从 stdin 读失败行，写出 xlsx（支持多工作表）
# --------------------------------------------------------------------------
def cmd_errrows(out_path):
    # 行协议：S<TAB>工作表名 开始一个新工作表（缺省名为 errrows）；
    #         H<TAB>... 表头；R<TAB>... 失败行（第一列「原行号」写成数值）
    groups = []          # [name, header, rows]
    cur = None
    for ln in sys.stdin.read().split("\n"):
        if not ln:
            continue
        parts = ln.split(TAB)
        tag = parts[0]
        if tag == "S":
            cur = [unesc(parts[1]) if len(parts) > 1 else "errrows", [], []]
            groups.append(cur)
            continue
        if cur is None:
            cur = ["errrows", [], []]
            groups.append(cur)
        if tag == "H":
            cur[1] = [unesc(x) for x in parts[1:]]
        elif tag == "R":
            vals = [unesc(x) for x in parts[1:]]
            # 第一列「原行号」写成数值，方便在 Excel 里排序/筛选
            if vals:
                try:
                    vals[0] = int(str(vals[0]).strip())
                except (TypeError, ValueError):
                    pass
            cur[2].append(vals)

    sheets = []
    for name, header, rows in groups:
        if not header and not rows:
            continue
        n_cols = max([len(header)] + [len(r) for r in rows]) if header or rows else 0
        sheets.append((name, [header] + rows, (1 + len(rows), n_cols)))
    write_xlsx(out_path, sheets)
    w("#OK\t%s" % os.path.abspath(out_path))
    return 0


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------
def main(argv):
    if not argv:
        note("helper: 缺少 mode 参数")
        return 2
    mode = argv[0]
    try:
        if mode == "probe":
            return cmd_probe(argv[1], argv[2],
                             int(argv[3]) if len(argv) > 3 else 0)
        if mode == "dump":
            return cmd_dump(argv[1], argv[2], int(argv[3]), int(argv[4]),
                            int(argv[5]), int(argv[6]), int(argv[7]),
                            argv[8] == "1", int(argv[9]),
                            int(argv[10]) if len(argv) > 10 else 0,
                            int(argv[11]) if len(argv) > 11 else 0)
        if mode == "errrows":
            return cmd_errrows(argv[1])
        if mode == "list":
            w("#SHEETS\t%s" % TAB.join(esc(x) for x in list_sheets(argv[1])))
            return 0
        note("helper: 未知 mode %r" % mode)
        return 2
    except HelperError as e:
        w("#FATAL\t%d\t%s" % (e.code, esc(str(e))))
        return e.code
    except ImportError as e:
        w("#FATAL\t%d\t%s" % (EXIT_ERROR,
                               esc("缺少 Python 依赖: %s。请先执行 pip install openpyxl xlrd" % e)))
        return EXIT_ERROR
    except BrokenPipeError:
        return 0
    except Exception as e:
        w("#FATAL\t%d\t%s" % (EXIT_ERROR, esc("%s: %s" % (type(e).__name__, e))))
        return EXIT_ERROR


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", newline="\n")
        sys.stderr.reconfigure(encoding="utf-8", newline="\n")
    except Exception:
        pass
    sys.exit(main(sys.argv[1:]))
