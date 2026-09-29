#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验 errrows.xlsx：表头第一列是「原行号」、最后一列是「失败原因」，且至少有一行数据。

    python3 tests/check_errrows.py <errrows.xlsx>
"""
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("缺少 openpyxl，无法校验", file=sys.stderr)
    sys.exit(1)


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    path = sys.argv[1]
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("errrows.xlsx 是空表", file=sys.stderr)
        return 1
    head = ["" if c is None else str(c) for c in rows[0]]
    if not head or head[0] != "原行号" or head[-1] != "失败原因":
        print("表头不符合预期: %r（应以「原行号」开头、「失败原因」结尾）" % head, file=sys.stderr)
        return 1
    data = [r for r in rows[1:] if r and r[0] is not None]
    if not data:
        print("errrows.xlsx 只有表头、没有失败行", file=sys.stderr)
        return 1
    for r in data:
        if not isinstance(r[0], int):
            print("第 %r 行的原行号不是整数" % (r[0],), file=sys.stderr)
            return 1
        if not r[-1]:
            print("第 %s 行缺少失败原因" % r[0], file=sys.stderr)
            return 1
    print("errrows.xlsx 校验通过：%d 列 × %d 条失败行，行号 %s"
          % (len(head), len(data), ", ".join(str(r[0]) for r in data)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
