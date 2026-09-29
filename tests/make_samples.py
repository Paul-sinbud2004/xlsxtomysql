#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成 xlsxtomysql 的测试样本文件（仅用于开发/自测，不属于工具本体）。

    python3 tests/make_samples.py [--big-rows 30000]

生成到 tests/samples/：
    01_normal.xlsx        标准表（类型齐全）
    02_messy.xlsx         非标准：空行 / 重复表头 / 分隔线 / 合计行
    03_merged_header.xlsx 两行表头 + 合并单元格
    04_dirty.xlsx         脏数据（少量非数值、超长文本）→ 用于验证 errrows.xlsx
    05_big.xlsx           大数据量（默认 30000 行）
    06_legacy.xls         Excel 2003 格式
    07_names.xlsx         字段名非法/超长/重名/为空/纯数字
    08_multi_sheet.xlsx   多工作表
"""
import argparse
import datetime as dt
import os
import random
import tempfile

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "samples")
# 沙箱里的 /tmp 往往只有 10MB tmpfs，openpyxl 流式写入会爆，改到样本目录
os.makedirs(OUT, exist_ok=True)
tempfile.tempdir = OUT

SURNAMES = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许"
GIVEN = ["梓涵", "浩然", "欣怡", "子轩", "雨泽", "诗琪", "俊杰", "可馨", "思远", "嘉怡"]


def name(i: int) -> str:
    return SURNAMES[i % len(SURNAMES)] + GIVEN[(i * 7) % len(GIVEN)]


def normal_rows(n: int):
    rnd = random.Random(20260929)
    for i in range(n):
        yield [
            name(i),
            8 + i % 8,                                             # 年龄 → tinyint
            round(120 + (i % 40) * 1.3, 1),                        # 身高 → decimal(x,1)
            dt.date(2018, 1, 1) + dt.timedelta(days=i % 3000),     # 出生日期 → date
            dt.time(7 + i % 2, (i * 3) % 60, 0),                   # 到校时间 → time
            f"第{i + 1}号同学的学习情况备注",                        # varchar
            round(12.5 + (i % 50) * 0.35, 2),                      # 午餐费 → decimal(x,2)
            bool(i % 2),                                           # 是否住校 → tinyint(1)
            f"138{i:08d}",                                         # 联系方式 → varchar(手机号)
            "该生在本学期表现良好，" * (2 + i % 3),                  # 长文本 → text
            f"2024-{(i % 12) + 1:02d}-{(i % 28) + 1:02d} 08:30:00",  # 字符串日期 → datetime
        ]


HEADERS = ["姓名", "年龄", "身高cm", "出生日期", "到校时间", "备注", "午餐费",
           "是否住校", "联系方式", "情况说明", "记录时间"]


def write_sheet(ws, rows, headers=HEADERS):
    ws.append(headers)
    for r in rows:
        ws.append(r)


def make_normal(n=500):
    wb = Workbook()
    ws = wb.active
    ws.title = "学生信息"
    write_sheet(ws, normal_rows(n))
    wb.save(os.path.join(OUT, "01_normal.xlsx"))
    return f"01_normal.xlsx  ({n} 行数据)"


def make_messy():
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    rows = list(normal_rows(40))
    out = [HEADERS]
    for i, r in enumerate(rows):
        out.append(r)
        if i == 9:
            out.append([None] * len(HEADERS))                 # 空行
        if i == 19:
            out.append(["—"] * len(HEADERS))                  # 分隔线
        if i == 24:
            out.append(HEADERS)                               # 重复表头
        if i == 34:
            out.append(["合计", 300, None, None, None, None, 999.99, None, None, None, None])
    for r in out:
        ws.append(r)
    wb.save(os.path.join(OUT, "02_messy.xlsx"))
    return "02_messy.xlsx  (含空行/分隔线/重复表头/合计行)"


def make_merged_header():
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["学生基本信息", None, None, None, "考勤", None, None])
    ws.append(["姓名", "年龄", "身高cm", "出生日期", "到校时间", "备注", "午餐费"])
    ws.merge_cells("A1:D1")
    ws.merge_cells("E1:G1")
    for i, r in enumerate(normal_rows(30)):
        ws.append(r[:7])
    wb.save(os.path.join(OUT, "03_merged_header.xlsx"))

    # 字段名行内部就有合并单元格的常见写法
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "数据"
    ws2.append(["姓名", "基本信息", None, "到校时间", "备注说明", None])
    ws2.merge_cells("B1:C1")
    ws2.merge_cells("E1:F1")
    for i, r in enumerate(normal_rows(30)):
        ws2.append([r[0], r[1], r[2], r[4], f"备注{i}", f"说明{i}"])
    wb2.save(os.path.join(OUT, "09_merged_single.xlsx"))
    return ("03_merged_header.xlsx (两行表头 + 合并，字段名行=2)\n"
            "         09_merged_single.xlsx (字段名行内合并单元格)")


def make_dirty(big=6000):
    """脏数据：把不合规的值放在抽样容易漏掉的位置。"""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("数据")
    ws.append(["姓名", "年龄", "身高cm", "出生日期", "到校时间", "备注", "午餐费"])
    for i, r in enumerate(normal_rows(big)):
        row = list(r[:7])
        if i >= big - 40 and (big - i) % 3 == 0:
            row[1] = "暂无"                     # 年龄列出现文本 → 整数列转换失败
            row[6] = "1.2345"                   # 午餐费列小数位超限 → decimal 失败
        ws.append(row)
    wb.save(os.path.join(OUT, "04_dirty.xlsx"))
    return f"04_dirty.xlsx  ({big} 行，末尾含脏数据)"


def make_big(n=30000):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("bigdata")
    ws.append(HEADERS)
    for r in normal_rows(n):
        ws.append(r)
    wb.save(os.path.join(OUT, "05_big.xlsx"))
    return f"05_big.xlsx  ({n} 行)"


def make_legacy():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("旧数据")
    style_date = xlwt.XFStyle()
    style_date.num_format_str = "YYYY-MM-DD"
    style_dt = xlwt.XFStyle()
    style_dt.num_format_str = "YYYY-MM-DD HH:MM:SS"
    for c, h in enumerate(["姓名", "年龄", "身高cm", "出生日期", "记录时间", "备注"]):
        ws.write(0, c, h)
    for i, r in enumerate(normal_rows(200)):
        ws.write(i + 1, 0, r[0])
        ws.write(i + 1, 1, r[1])
        ws.write(i + 1, 2, r[2])
        ws.write(i + 1, 3, dt.datetime.combine(r[3], dt.time()), style_date)
        ws.write(i + 1, 4, dt.datetime.combine(r[3], r[4]), style_dt)
        ws.write(i + 1, 5, r[5])
    wb.save(os.path.join(OUT, "06_legacy.xls"))
    return "06_legacy.xls     (Excel 2003 格式，200 行)"


def make_names():
    wb = Workbook()
    ws = wb.active
    ws.title = "怪字段名"
    ws.append(["姓名", "姓 名", "姓名", "", "2024", "这个字段名字特别特别长" * 6,
               "a-b(c)", "金额(元)", "备注"])
    for i in range(10):
        ws.append([f"张三{i}", "x", "y", "空字段名的值", i, "很长" * 5, "a", 1.5, "备注"])
    wb.save(os.path.join(OUT, "07_names.xlsx"))
    return "07_names.xlsx  (非法字符/超长/重名/空名/纯数字字段名)"


def make_xls_merge():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("旧表")
    ws.write_merge(0, 0, 1, 2, "基本信息")
    ws.write(0, 0, "姓名")
    ws.write(0, 3, "备注说明")
    ws.write(0, 4, "补充说明")
    for i in range(20):
        ws.write(i + 1, 0, f"学生{i}")
        ws.write(i + 1, 1, 10 + i)
        ws.write(i + 1, 2, 160.5 + i)
        ws.write(i + 1, 3, f"备注{i}")
        ws.write(i + 1, 4, f"说明{i}")
    wb.save(os.path.join(OUT, "11_xls_merge.xls"))
    return "11_xls_merge.xls  (.xls 字段名行内含合并单元格)"


def make_empty():
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws.append(["姓名", "分数", "空列", "备注"])
    ws.append(["甲", 90, None, "合格"])
    ws.append(["乙", None, None, None])
    ws.append(["丙", 88, None, ""])
    wb.save(os.path.join(OUT, "12_empty.xlsx"))
    return "12_empty.xlsx  (空单元格 / 空列)"


def make_multi_sheet():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "第一张表"
    write_sheet(ws1, list(normal_rows(20)))
    ws2 = wb.create_sheet("第二张表")
    ws2.append(["编号", "商品", "单价", "数量"])
    for i in range(20):
        ws2.append([i + 1, f"商品{i}", round(9.9 + i, 2), i % 7 + 1])
    ws3 = wb.create_sheet("说明")
    ws3.append(["本文件用于测试多工作表"])
    wb.save(os.path.join(OUT, "08_multi_sheet.xlsx"))
    return "08_multi_sheet.xlsx  (3 个工作表)"


def make_edge_cases():
    """特殊表格边界：日期写法 / 合并填充 / 中间空行 / 隐藏列 / 百分比 / 时区。

    对应探针实测里确认过的一批行为，固化成回归样本。
    """
    # 13 文本日期：ISO T 分隔、紧凑、斜杠、毫秒、上下午、超 24 小时时长
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "登记时间", "出生日期", "加班时长"])
    rows13 = [
        ("甲", "2024-01-01T08:30:00", "20240102", "08:30:00"),
        ("乙", "2024/02/03 09:15:00", "1998/06/07", "25:30:00"),
        ("丙", "2024-03-04T10:20:30.500", "2000-12-31", "100:00:00"),
        ("丁", "2024-04-05 11:00 AM", "1985.09.10", "00:45:00"),
    ]
    for r in rows13:
        ws.append(list(r))
    wb.save(os.path.join(OUT, "13_date_text.xlsx"))

    # 14 数据区纵向合并：一个学生占 3 行，只有第一行有名字
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "科目", "分数"])
    data14 = [("张三", "语文", 90), (None, "数学", 85), (None, "英语", 78),
              ("李四", "语文", 88), (None, "数学", 92), (None, "英语", 81)]
    for r in data14:
        ws.append(list(r))
    ws.merge_cells("A2:A4")
    ws.merge_cells("A5:A7")
    wb.save(os.path.join(OUT, "14_merge_vertical.xlsx"))

    # 15 字段名行落在纵向合并区里：字段名在合并区左上角的第 1 行，
    #    第 2 行（字段名行参数指向它）读出来是空的，须回读上一行
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "语文", "数学"])
    ws.append([None, None, None])
    for i in range(6):
        ws.append([name(i), 80 + i, 70 + i])
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:C2")
    wb.save(os.path.join(OUT, "15_header_vmerge.xlsx"))

    # 16 数据区中间空行：应被判定为非标准格式并阻断（exit=2）
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "分数"])
    ws.append(["甲", 90])
    ws.append([None, None])
    ws.append(["乙", 85])
    wb.save(os.path.join(OUT, "16_blank_mid.xlsx"))

    # 17 隐藏列 + 百分比格式
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "完成率", "内部备注", "折扣"])
    for i in range(6):
        ws.append([name(i), 0.125 + i * 0.01, f"内部{i}", 0.85])
        ws[f"B{i + 2}"].number_format = "0.00%"
    ws.column_dimensions["C"].hidden = True
    wb.save(os.path.join(OUT, "17_hidden_pct.xlsx"))

    # 18 带时区偏移的日期时间：偏移会被丢弃，应给出提示
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "上传时间"])
    for i in range(5):
        ws.append([name(i), f"2024-01-0{i + 1}T08:30:00+08:00"])
    wb.save(os.path.join(OUT, "18_tz_offset.xlsx"))

    return ("13_date_text.xlsx      (文本日期：ISO T / 紧凑 / 斜杠 / 毫秒 / 上下午 / 超 24h)\n"
            "         14_merge_vertical.xlsx (数据区纵向合并：一人占 3 行)\n"
            "         15_header_vmerge.xlsx  (字段名行落在纵向合并区内)\n"
            "         16_blank_mid.xlsx      (数据区中间空行，应阻断 exit=2)\n"
            "         17_hidden_pct.xlsx     (隐藏列 + 百分比格式)\n"
            "         18_tz_offset.xlsx      (带时区偏移的日期时间)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--big-rows", type=int, default=30000)
    ap.add_argument("--normal-rows", type=int, default=500)
    args = ap.parse_args()
    os.makedirs(OUT, exist_ok=True)
    makers = [
        lambda: make_normal(args.normal_rows),
        make_messy,
        make_merged_header,
        make_dirty,
        lambda: make_big(args.big_rows),
        make_legacy,
        make_names,
        make_multi_sheet,
        make_xls_merge,
        make_empty,
        make_edge_cases,
    ]
    for m in makers:
        try:
            print("  生成", m())
        except ImportError as e:
            print("  跳过（缺少依赖）:", e)


if __name__ == "__main__":
    main()
